import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from dotenv import load_dotenv
from openai import OpenAI

load_dotenv(".env")
client = OpenAI()

DATA_DIR = Path("data")

SOURCE_DATES = {
    "01_Project_Overview_NexusCert.docx": "2026-01-06",
    "02_Meeting_Minutes_Architecture_18-May-2026.docx": "2026-05-18",
    "03_Meeting_Minutes_UAT_Phase2_24-Jun-2026.docx": "2026-06-24",
    "04_WSR_08-Jun-2026_to_12-Jun-2026.docx": "2026-06-12",
    "05_WSR_22-Jun-2026_to_26-Jun-2026.docx": "2026-06-26",
    "06_Project_Control_Tracker_NexusCert.xlsx": "2026-06-26",
}


def read_docx(path):
    doc = Document(path)
    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip():
                parts.append(text)

    return "\n".join(parts)


def read_xlsx(path):
    workbook = load_workbook(path, data_only=True)
    parts = []

    for sheet in workbook.worksheets:
        parts.append(f"Sheet: {sheet.title}")

        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts)


def chunk_text(text, chunk_size=1200, overlap=200):
    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])

        if end >= len(text):
            break

        start = end - overlap

    return chunks


records = []

for path in sorted(DATA_DIR.iterdir()):

    if path.suffix.lower() == ".docx":
        text = read_docx(path)

    elif path.suffix.lower() == ".xlsx":
        text = read_xlsx(path)

    else:
        continue

    chunks = chunk_text(text)

    for number, chunk in enumerate(chunks, start=1):
        records.append({
            "source": path.name,
            "source_date": SOURCE_DATES.get(path.name, "1900-01-01"),
            "chunk": number,
            "text": chunk
        })


print(f"Preparing {len(records)} chunks...")

response = client.embeddings.create(
    model="text-embedding-3-small",
    input=[record["text"] for record in records]
)

for record, embedding in zip(records, response.data):
    record["embedding"] = embedding.embedding


with open("index.json", "w", encoding="utf-8") as f:
    json.dump(records, f)


print(f"Created searchable index with {len(records)} chunks.")
print("Saved as index.json")
