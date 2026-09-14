from pathlib import Path
from docx import Document
from openpyxl import load_workbook

DATA_DIR = Path("data")

def read_docx(path):
    doc = Document(path)
    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)

def read_xlsx(path):
    workbook = load_workbook(path, data_only=True)
    parts = []

    for sheet in workbook.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts)

for path in sorted(DATA_DIR.iterdir()):
    if path.suffix.lower() == ".docx":
        text = read_docx(path)
    elif path.suffix.lower() == ".xlsx":
        text = read_xlsx(path)
    else:
        continue

    print(f"{path.name}: {len(text):,} characters extracted")
